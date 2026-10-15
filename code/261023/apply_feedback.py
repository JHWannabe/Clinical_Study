# 261023.pptx에 미팅 피드백 반영: Sens90 슬라이드를 부록으로, ROC·BA·가설 비교·Next steps 슬라이드 추가, BH q 병기, 불균형 지표 추가
import copy, re, sys
from pathlib import Path
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Emu
sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]; OUT = ROOT / "outputs/261023"; PPT = ROOT / "docs/261023.pptx"
prs = Presentation(PPT); S = list(prs.slides)
DIS = ["HTN", "DM", "CKD"]
hyp = pd.read_excel(OUT / "hypothesis_ba.xlsx"); prev = pd.read_excel(OUT / "prevalence.xlsx").set_index("disease")
NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


# 문단 para의 첫 run에 text를 넣고 나머지 run은 비움
def set_para(para, text):
    for i, r in enumerate(para.runs): r.text = text if i == 0 else ""


# 셀의 첫 문단만 남기고 text를 넣음
def set_cell(cell, text):
    tf = cell.text_frame
    for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
    set_para(tf.paragraphs[0], text)


# 슬라이드의 이름 접두어별 도형
shp = lambda s, name: next(x for x in s.shapes if x.name.startswith(name))


# 본문 placeholder의 문단을 texts로 교체 (필요하면 마지막 문단을 복제)
def set_body(s, texts):
    ps = shp(s, "Content").text_frame.paragraphs
    while len(ps) < len(texts):
        ps[-1]._p.addnext(copy.deepcopy(ps[-1]._p)); ps = shp(s, "Content").text_frame.paragraphs
    for p in ps[len(texts):]: p._p.getparent().remove(p._p)
    for p, t in zip(shp(s, "Content").text_frame.paragraphs, texts): set_para(p, t)


# 템플릿 슬라이드 tpl의 도형(그림 제외)을 복사한 새 슬라이드
def clone(tpl, title, badge, texts):
    s = prs.slides.add_slide(tpl.slide_layout)
    for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
    for sh in tpl.shapes:
        if sh.shape_type != 13: s.shapes._spTree.append(copy.deepcopy(sh._element))
    set_para(shp(s, "Title").text_frame.paragraphs[0], title); set_para(shp(s, "Rounded").text_frame.paragraphs[0], badge)
    set_body(s, texts)
    return s


# BH-FDR q
def bh(p):
    n = len(p); order = sorted(range(n), key=lambda i: p[i]); q = [0] * n; m = 1
    for r in range(n, 0, -1):
        i = order[r - 1]; m = min(m, p[i] * n / r); q[i] = m
    return q


base = hyp[hyp.feature_set == "clinic4"].set_index("disease")
bst = hyp[hyp.feature_set.str.startswith("best")].set_index("disease")

# 1) 슬라이드 2 Summary
set_body(S[1], [
    "목적: 기본 정보(나이·성별·키·몸무게)에 CT 몸 구성을 추가했을 때 HTN·DM·CKD 예측 변화",
    "자료: 강남·신촌 합친 5-fold CV(3회 반복), 질환군 vs 정상군",
    "기본 정보만: AUC HTN 0.77, DM 0.77, CKD 0.75 (모두 정상 예측 시 정확도 72%·82%·94%)",
    "CT 몸 구성 추가: DM 0.80 (+0.036), CKD 0.81 (+0.060), HTN +0.012 (p=0.149)",
    "가설 비교: DM은 총량만으로 +0.025, 분포(SD·첨도) 추가 이득 없음",
    "한계: 구간·조합 사후 탐색, 다중비교 보정(BH) 필요, CKD 환자 수 적음, 단면 자료"])

# 2) 슬라이드 8 baseline: Sens/Spec 대신 불균형 지표
s = S[7]
set_body(s, ["모델: clinic4 (Age·Sex·Height·Weight 4개 입력) 로지스틱 회귀, 규제 강도는 train fold 안에서 선택",
             "평가: 합친 5-fold CV (3회 반복), AUC는 mean ± SD, BA는 BA 최대 cutoff 기준"])
t = shp(s, "Table").table
rows = {1: ("환자 수 (질환자)", lambda d: f"{prev.n[d]:,} ({prev.n_pos[d]:,})"), 2: ("질환자 비율", lambda d: f"{prev.prevalence[d]:.1%}"),
        3: ("모두 정상 예측 시 정확도", lambda d: f"{prev.majority_acc[d]:.1%}"), 4: ("Balanced accuracy", lambda d: f"{base.BA[d]:.3f}")}
for r, (lab, fn) in rows.items():
    set_cell(t.cell(r, 0), lab)
    for c, d in enumerate(DIS, 1): set_cell(t.cell(r, c), fn(d))

# 3) 슬라이드 9 사례 표에 BH q 병기 (질환별 7개 사례, p<0.001은 0.001로 두어 q는 상한)
t = shp(S[8], "Table").table
for c in (1, 2, 3):
    ms = [re.search(r"p([=<])([0-9.]+)", t.cell(r, c).text) for r in range(1, 8)]
    q = bh([float(m.group(2)) for m in ms])
    for r, (m, qq) in enumerate(zip(ms, q), 1):
        set_cell(t.cell(r, c), t.cell(r, c).text.replace(")", f", q{'<' if m.group(1) == '<' else '='}{qq:.3f})"))

# 4) 새 슬라이드: ROC + BA, 가설 비교, Next steps
tpl12, tpl14 = S[11], S[13]
roc = clone(tpl12, "Final best model: ROC curve and balanced accuracy", "Results", [
    "곡선이 왼쪽 위에 붙을수록 좋음, 대각선은 무작위 (합친 5-fold CV OOF 예측)"])
shp(roc, "Content").height = Emu(500000)
w = 8400000; im = Image.open(OUT / "roc_baseline_vs_best.png"); h = int(w * im.height / im.width)
roc.shapes.add_picture(str(OUT / "roc_baseline_vs_best.png"), Emu((12192000 - w) // 2), Emu(1850000), Emu(w), Emu(h))
tb = shp(roc, "Table"); tb.top = Emu(1850000 + h + 60000); t = tb.table
for r in t.rows: r.height = Emu(380000)
tb.height = Emu(1520000)
for c, x in enumerate(["Disease", "AUC (clinic4)", "AUC (best)", "Balanced acc. (clinic4)", "Balanced acc. (best)"]): set_cell(t.cell(0, c), x)
for r, d in enumerate(DIS, 1):
    for c, x in enumerate([d, f"{base.AUC[d]:.3f} ± {base.AUC_sd[d]:.3f}", f"{bst.AUC[d]:.3f} ± {bst.AUC_sd[d]:.3f}", f"{base.BA[d]:.3f}", f"{bst.BA[d]:.3f}"]): set_cell(t.cell(r, c), x)

hs = clone(tpl12, "Hypothesis: total vs. distribution", "Results", [
    "가설: 조직 분포(SD·첨도)가 총량(구간 합산)보다 의미 있다. AUC (ΔAUC, p)"])
t = shp(hs, "Table").table
names = ["clinic4", "+ 총량 (구간 합산 16)", "+ 총량 + 분포 (SD·첨도 포함 48)", "best (구간 평균·합산 32)"]
for c, x in enumerate(["Disease", "clinic4", "+ 총량 (합산 16)", "+ 총량 + 분포 (48)", "best (평균·합산 32)"]): set_cell(t.cell(0, c), x)
for r, d in enumerate(DIS, 1):
    set_cell(t.cell(r, 0), d)
    for c, n in enumerate(names, 1):
        g = hyp[(hyp.disease == d) & (hyp.feature_set == n)].iloc[0]
        txt = f"{g.AUC:.3f} ± {g.AUC_sd:.3f}"
        if n != "clinic4": txt += f" ({g.dAUC:+.3f}, p{'<0.001' if g.p_cv < .001 else f'={g.p_cv:.3f}'})"
        set_cell(t.cell(r, c), txt)

ns = clone(tpl14, "Next steps (meeting feedback)", "Next steps", [
    "지표: Sens90 대신 AUC·ROC·balanced accuracy로 보고, 다음 발표에 코호트별 ROC 추가",
    "가설: 구간 정의의 임상 근거(문헌) 확보 후 사전 지정, 사례 547개는 BH-FDR q로 보고, 참석자 3과 가설·그래프 분석 협의",
    "코호트: 두 병원 합산 대신 분리 유지, 병원 간 차이가 큰 VAT 중심으로 정리",
    "불균형: DM AUC 0.77 원인 확인(유병률·누수), 정상:질환 3:7 이상으로 언더·오버샘플링과 피처 단 SMOTE 비교",
    "PSM: 질환군과 정상군 매칭 후 AUC(기존 수행분 약 0.8) 수치 재확인 후 공유, 질환 비율 유지 vs 매칭 중 기준 정리",
    "영상 모델: 12장 입력은 2.5D(채널 적층) ResNet, 환자 단위 판정은 장별 예측 합산도 비교, 선행 연구(이태희 연구원, 임재열 교수 과제) 참고"])

# 5) Limitations (슬라이드 14)
set_body(S[13], [
    "CKD 환자 수 적음 (강남 45명, 신촌 63명), 점수 변동 큼 (best AUC SD 0.052)",
    "구간·사례 547개를 사전 가설 없이 탐색, 다중비교(BH) 보정 전 p값은 신뢰 어려움",
    "best case 선택과 성능 평가에 같은 CV 사용, AUC 과대 가능성",
    "독립된 외부 test set 없이 합친 5-fold CV만 평가",
    "병원 간 차이: VAT 분포 차이 큼, TAMA 차이는 강남에서만 유의",
    "질환 불균형(질환자 6~28%): 모두 정상 예측해도 정확도 72~94%",
    "정상군은 HTN·DM·CKD 모두 없는 환자, 단면 자료로 인과 해석 불가"])

# 6) Sens90 슬라이드(13)는 부록 끝으로, 제목 변경
set_para(shp(S[12], "Title").text_frame.paragraphs[0], "Appendix D1. Specificity at 90% sensitivity")
set_para(shp(S[12], "Rounded").text_frame.paragraphs[0], "Appendix")

set_body(S[12], ["민감도 90% 이상을 만족하는 가장 큰 cutoff (Sens90)에서 Spec 비교, 값은 mean ± SD", "Sens는 baseline 값만 표기 (best model과 차이 0.001 이하)"])
# 7) 순서: 1-12, ROC, 가설, Limitations, Next steps, 종료 표지, 부록(16..), Sens90
lst = prs.slides._sldIdLst; ids = list(lst); new = ids[len(S):]
order = ids[:12] + [new[0], new[1], ids[13], new[2], ids[14]] + ids[15:len(S)] + [ids[12]]
for e in ids: lst.remove(e)
for e in order: lst.append(e)

# 슬라이드 번호 필드 갱신
for i, s in enumerate(prs.slides, 1):
    for f in s._element.iter(NS + "fld"):
        for tt in f.iter(NS + "t"): tt.text = str(i)
prs.save(PPT); print(len(prs.slides), "slides saved")
