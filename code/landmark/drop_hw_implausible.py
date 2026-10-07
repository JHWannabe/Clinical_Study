# data/*_landmark_filtered.xlsx의 patient_metadata 시트에서 hw_implausible 칼럼 삭제 (이미 없으면 건너뜀, 재실행 안전)
import openpyxl

for c in ["gangnam", "sinchon"]:
    p = f"data/{c}_landmark_filtered.xlsx"
    wb = openpyxl.load_workbook(p)
    ws = wb["patient_metadata"]
    names = [x.value for x in ws[1]]
    if "hw_implausible" in names:
        ws.delete_cols(names.index("hw_implausible") + 1)
        wb.save(p)
    print(c, "done", [x.value for x in ws[1]])
